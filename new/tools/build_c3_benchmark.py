"""Build C3's individual-specialty code panel from the supplied workbook."""
import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def build(source: Path, output: Path) -> int:
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["Spec-Proc Weights"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value).strip() for value in next(rows)]
    code_i, weight_i = header.index("PROC_CD"), header.index("SPEC_PROC_WEIGHT")
    weights = defaultdict(float)
    for row in rows:
        code = str(row[code_i] or "").strip().upper()
        if code and float(row[weight_i]) > 0:
            weights[code] += float(row[weight_i])
    workbook.close()
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["billing_code", "weight"])
        writer.writerows((code, str(int(weight)) if weight.is_integer() else repr(weight))
                         for code, weight in sorted(weights.items()))
    return len(weights)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(f"Wrote {build(args.source, args.output)} distinct codes to {args.output}")
