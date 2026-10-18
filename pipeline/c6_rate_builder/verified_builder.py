"""Conservative C5 -> C6 projection with no synthetic addresses or silent rate collapse."""
from __future__ import annotations

import argparse
import csv
import json
import sqlite3
from collections import Counter
from pathlib import Path

from openpyxl import Workbook

RATE_COLUMNS = ("NPI", "TIN", "Address", "Billing Code", "Specialty", "Rate")
AUDIT_COLUMNS = RATE_COLUMNS + ("Tin Type", "Source File ID", "Source Locator", "Rate Context", "C5 Evidence IDs")


def build_verified_rates(selected_rates: list[Path], c5_locations: Path, output_dir: Path,
                         *, allowed_plan_attributions: set[str] | None = None) -> dict:
    """Export only exact PPO/TIN/site matches from rate files with accepted plan attribution.

    The default empty allowlist intentionally withholds rates from mixed MRFs.
    A projected six-column key with multiple C4 contexts is sent to an ambiguity
    file rather than choosing an unrelated rate or losing a modifier/setting.
    """
    allowed = allowed_plan_attributions or set()
    output_dir.mkdir(parents=True, exist_ok=True)
    scratch = output_dir / "c6_work.sqlite"
    if scratch.exists():
        scratch.unlink()
    db = sqlite3.connect(scratch)
    counts = Counter()
    try:
        db.executescript("""
            CREATE TABLE sites (npi TEXT, tin_type TEXT, tin TEXT, address_key TEXT,
                address TEXT, specialty TEXT, evidence_ids TEXT,
                PRIMARY KEY(npi,tin_type,tin,address_key));
            CREATE INDEX sites_pair ON sites(npi,tin_type,tin);
            CREATE TABLE projected (npi TEXT, tin_type TEXT, tin TEXT, address TEXT,
                address_key TEXT, code TEXT, specialty TEXT, rate TEXT,
                source_file_id TEXT, source_locator TEXT, context TEXT, evidence_ids TEXT);
            CREATE INDEX projected_key ON projected(npi,tin_type,tin,address_key,code);
        """)
        with c5_locations.open(newline="", encoding="utf-8-sig") as stream:
            for site in csv.DictReader(stream):
                if not (site["practice_site_status"] == "corroborated" and
                        site["tin_link_status"] == "proven" and
                        site["ppo_participation_status"] == "confirmed"):
                    continue
                db.execute("INSERT OR IGNORE INTO sites VALUES (?,?,?,?,?,?,?)",
                           (site["npi"], site["tin_type"], site["tin"], site["address_key"],
                            site["address"], site["specialty"], site["source_ids"]))
        db.commit()
        for path in selected_rates:
            with path.open(newline="", encoding="utf-8-sig") as stream:
                for rate in csv.DictReader(stream):
                    counts["selected_c4_rows"] += 1
                    if rate.get("plan_attribution", "") not in allowed or rate.get("file_classification") == "MIXED":
                        counts["withheld_unattributed_rate_rows"] += 1
                        continue
                    sites = db.execute("SELECT address_key,address,specialty,evidence_ids FROM sites WHERE npi=? AND tin_type=? AND tin=?",
                                       (rate.get("npi", ""), rate.get("tin_type", ""), rate.get("tin", ""))).fetchall()
                    if not sites:
                        counts["withheld_no_verified_site_rows"] += 1
                        continue
                    context = {key: rate.get(key, "") for key in (
                        "billing_code_type", "billing_code_type_version", "billing_class", "setting",
                        "service_code", "billing_code_modifiers", "negotiated_type", "rate_unit",
                        "negotiation_arrangement", "expiration_date")}
                    for address_key, address, specialty, evidence_ids in sites:
                        db.execute("INSERT INTO projected VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                                   (rate["npi"], rate["tin_type"], rate["tin"], address, address_key,
                                    rate["billing_code"], specialty, rate["rate"],
                                    rate.get("source_file_id", ""), rate.get("source_locator", ""),
                                    json.dumps(context, sort_keys=True), evidence_ids))
                        counts["site_expanded_rate_contexts"] += 1
            db.commit()
        audit_path = output_dir / "c6_rate_context_audit.csv"
        ambiguity_path = output_dir / "c6_ambiguous_projection.csv"
        csv_path = output_dir / "c6_verified_rates.csv"
        xlsx_path = output_dir / "Centene_CA_Commercial_PPO_Provider_Rates.xlsx"
        book = Workbook(write_only=True)
        sheet_num = 1
        sheet = book.create_sheet(f"Provider_Rates_{sheet_num}")
        sheet.append(RATE_COLUMNS)
        sheet_rows = 0
        with audit_path.open("w", newline="", encoding="utf-8") as audit_stream, \
             ambiguity_path.open("w", newline="", encoding="utf-8") as ambiguity_stream, \
             csv_path.open("w", newline="", encoding="utf-8") as csv_stream:
            audit = csv.writer(audit_stream)
            ambiguous = csv.writer(ambiguity_stream)
            final = csv.writer(csv_stream)
            audit.writerow(AUDIT_COLUMNS)
            ambiguous.writerow(AUDIT_COLUMNS)
            final.writerow(RATE_COLUMNS)
            query = """SELECT p.npi,p.tin_type,p.tin,p.address,p.code,p.specialty,p.rate,
                       p.source_file_id,p.source_locator,p.context,p.evidence_ids,
                       COUNT(*) OVER (PARTITION BY p.npi,p.tin_type,p.tin,p.address_key,p.code)
                       FROM projected p ORDER BY p.npi,p.tin_type,p.tin,p.address_key,p.code,p.rowid"""
            for npi, kind, tin, address, code, specialty, amount, file_id, locator, context, evidence_ids, group_size in db.execute(query):
                rate_row = (npi, tin, address, code, specialty, amount)
                audit_row = rate_row + (kind, file_id, locator, context, evidence_ids)
                audit.writerow(audit_row)
                if group_size > 1:
                    ambiguous.writerow(audit_row)
                    counts["ambiguous_context_rows"] += 1
                    continue
                final.writerow(rate_row)
                if sheet_rows == 1_000_000:
                    sheet_num += 1
                    sheet = book.create_sheet(f"Provider_Rates_{sheet_num}")
                    sheet.append(RATE_COLUMNS)
                    sheet_rows = 0
                sheet.append(rate_row)
                sheet_rows += 1
                counts["exported_rate_rows"] += 1
        book.save(xlsx_path)
        summary = {"counts": dict(counts), "allowed_plan_attributions": sorted(allowed),
                   "outputs": [p.name for p in (csv_path, xlsx_path, audit_path, ambiguity_path)],
                   "rule": "Requires corroborated site, direct TIN/site evidence, exact PPO directory hit, and PPO-only rate attribution; ambiguous context projections withheld."}
        (output_dir / "c6_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
        return summary
    finally:
        db.close()


def main() -> None:
    from pipeline.c5_address_mapper.resolve import selected_paths

    parser = argparse.ArgumentParser(description="Project attributed C4 rates onto verified C5 PPO practice sites")
    parser.add_argument("--c4-manifest", type=Path)
    parser.add_argument("--selected-rates", type=Path, action="append", default=[])
    parser.add_argument("--c5-locations", type=Path, required=True)
    parser.add_argument("--allow-plan-attribution", action="append", default=[])
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    if not args.c5_locations.is_file():
        raise FileNotFoundError(args.c5_locations)
    summary = build_verified_rates(
        selected_paths(args.c4_manifest, args.selected_rates), args.c5_locations,
        args.output_dir, allowed_plan_attributions=set(args.allow_plan_attribution))
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
