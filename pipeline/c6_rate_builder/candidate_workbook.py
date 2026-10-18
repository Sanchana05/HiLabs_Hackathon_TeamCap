"""Export the C6 research candidates with their evidence and scope in one workbook."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def _csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream)
        return list(reader.fieldnames or []), list(reader)


def write_candidate_workbook(c6_dir: Path, destination: Path) -> Path:
    """Write the six-column result and an explicit row-level evidence trail."""
    columns, candidates = _csv_rows(c6_dir / "c6_supported_rate_candidates.csv")
    expected = ["NPI", "TIN", "Address", "Billing Code", "Specialty", "Rate"]
    if columns != expected:
        raise ValueError(f"Unexpected candidate columns: {columns}")
    evidence_columns, evidence = _csv_rows(c6_dir / "c6_candidate_evidence.csv")
    if len(candidates) != len(evidence):
        raise ValueError("Candidate and evidence row counts differ")
    book = Workbook()
    rate_sheet = book.active
    rate_sheet.title = "Rate candidates"
    rate_sheet.append(expected)
    for row in candidates:
        rate_sheet.append([row[name] if name != "Rate" else float(row[name]) for name in expected])
    evidence_sheet = book.create_sheet("Evidence")
    evidence_sheet.append(evidence_columns)
    for row in evidence:
        evidence_sheet.append([row[name] for name in evidence_columns])
    scope = book.create_sheet("Scope and sources")
    for line in (
        ("Status", "RESEARCH CANDIDATES — NOT VERIFIED PPO PROVIDER RATES"),
        ("Source", "C4 selected published rates joined to C5 corroborated practice sites"),
        ("PPO limitation", "The selected MRF is shared with non-PPO plans; PPO-specific rate attribution is unproven."),
        ("TIN/site limitation", "A direct TIN-to-practice-site roster is required."),
        ("PPO/site limitation", "An exact plan-filtered PPO provider directory is required."),
        ("Audit", "See the Evidence sheet and c6_candidate_evidence.csv for each row."),
    ):
        scope.append(line)
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions if sheet != scope else "A1:B6"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            sheet.column_dimensions[letter].width = min(60, max(14, max(len(str(c.value or "")) for c in column[:250]) + 2))
    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)
    return destination
