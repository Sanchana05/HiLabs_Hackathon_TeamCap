import csv

from openpyxl import load_workbook

from pipeline.c6_rate_builder.candidate_workbook import write_candidate_workbook


def test_candidate_workbook_contains_six_columns_and_evidence(tmp_path):
    c6_dir = tmp_path / "c6"
    c6_dir.mkdir()
    fields = ["NPI", "TIN", "Address", "Billing Code", "Specialty", "Rate"]
    row = dict(zip(fields, ["0012345678", "000123456", "10 MAIN ST, CA", "99213", "Family Medicine", "100.25"]))
    with (c6_dir / "c6_supported_rate_candidates.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerow(row)
    with (c6_dir / "c6_candidate_evidence.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields + ["Decision"])
        writer.writeheader()
        writer.writerow({**row, "Decision": "RESEARCH_CANDIDATE_NOT_DIRECTORY_VERIFIED"})
    path = write_candidate_workbook(c6_dir, tmp_path / "candidates.xlsx")
    book = load_workbook(path, read_only=True, data_only=True)
    assert book.sheetnames == ["Rate candidates", "Evidence", "Scope and sources"]
    assert list(book["Rate candidates"].values) == [tuple(fields),
        ("0012345678", "000123456", "10 MAIN ST, CA", "99213", "Family Medicine", 100.25)]
    assert book["Evidence"]["G2"].value == "RESEARCH_CANDIDATE_NOT_DIRECTORY_VERIFIED"
    book.close()
